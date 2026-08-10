from __future__ import annotations

import argparse
import hashlib
import json
import math
import mimetypes
import re
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
WELLS_PATH = ROOT / "data" / "wells.json"
ATTACHMENTS_DIR = ROOT / "data" / "attachments"
WATER_RIGHT_RE = re.compile(r"[BK]\d{7}")


def clean_text(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def clean_water_right(value):
    match = WATER_RIGHT_RE.search(str(value or ""))
    return match.group(0) if match else None


def clean_number(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"-?\d+(?:\.\d+)?", str(value).replace(",", ""))
    return float(match.group(0)) if match else None


def compact_number(value):
    if value is None:
        return None
    return int(value) if float(value).is_integer() else float(value)


def roc_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return f"{value.year - 1911:03d}.{value.month:02d}.{value.day:02d}"
    digits = re.sub(r"\D", "", str(value))
    if len(digits) == 7:
        return f"{int(digits[:3]):03d}.{digits[3:5]}.{digits[5:7]}"
    text = clean_text(value)
    return text.replace("/", ".") if text else None


def twd97_to_wgs84(x, y):
    a = 6378137.0
    b = 6356752.314245
    longitude0 = math.radians(121)
    k0 = 0.9999
    dx = 250000
    e = math.sqrt(1 - (b * b) / (a * a))
    x -= dx
    m = y / k0
    mu = m / (a * (1 - e**2 / 4 - 3 * e**4 / 64 - 5 * e**6 / 256))
    e1 = (1 - math.sqrt(1 - e**2)) / (1 + math.sqrt(1 - e**2))
    j1 = 3 * e1 / 2 - 27 * e1**3 / 32
    j2 = 21 * e1**2 / 16 - 55 * e1**4 / 32
    j3 = 151 * e1**3 / 96
    j4 = 1097 * e1**4 / 512
    fp = mu + j1 * math.sin(2 * mu) + j2 * math.sin(4 * mu) + j3 * math.sin(6 * mu) + j4 * math.sin(8 * mu)
    e2 = (e * a / b) ** 2
    c1 = e2 * math.cos(fp) ** 2
    t1 = math.tan(fp) ** 2
    r1 = a * (1 - e**2) / (1 - e**2 * math.sin(fp) ** 2) ** 1.5
    n1 = a / math.sqrt(1 - e**2 * math.sin(fp) ** 2)
    d = x / (n1 * k0)
    latitude = fp - (n1 * math.tan(fp) / r1) * (
        d**2 / 2
        - (5 + 3 * t1 + 10 * c1 - 4 * c1**2 - 9 * e2) * d**4 / 24
        + (61 + 90 * t1 + 298 * c1 + 45 * t1**2 - 252 * e2 - 3 * c1**2) * d**6 / 720
    )
    longitude = longitude0 + (
        d
        - (1 + 2 * t1 + c1) * d**3 / 6
        + (5 - 2 * c1 + 28 * t1 - 3 * c1**2 + 8 * e2 + 24 * t1**2) * d**5 / 120
    ) / math.cos(fp)
    return round(math.degrees(latitude), 7), round(math.degrees(longitude), 7)


def is_valid_twd97(x, y):
    return x is not None and y is not None and 100000 <= x <= 400000 and 2400000 <= y <= 2900000


def image_extension(image):
    return {"jpeg": ".jpg", "jpg": ".jpg", "png": ".png", "gif": ".gif"}.get(
        str(image.format or "").lower(), ".jpg"
    )


def parse_workbook(path):
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    total = workbook["總表"]
    registry = []
    row = 2
    while isinstance(total.cell(row, 1).value, (int, float)):
        values = [total.cell(row, column).value for column in range(1, 27)]
        registry.append(
            {
                "recordNo": int(values[0]),
                "station": clean_text(values[1]),
                "name": clean_text(values[2]),
                "waterRightNo": clean_water_right(values[3]),
                "electricityNo": clean_text(values[4]),
                "agriculturalPower": clean_text(values[5]),
                "address": clean_text(values[6]),
                "depthMeters": compact_number(clean_number(values[7])),
                "diameterMm": compact_number((clean_number(values[8]) or 0) * 1000) if clean_number(values[8]) is not None else None,
                "pumpHorsepower": compact_number(clean_number(values[9])),
                "pumpOutletInch": compact_number(clean_number(values[10])),
                "planFlowCms": compact_number(clean_number(values[11])),
                "benefitedAreaHa": compact_number(clean_number(values[12])),
                "registeredFlowCms": compact_number(clean_number(values[13])),
                "irrigationSystem": clean_text(values[14]),
                "purpose": clean_text(values[15]),
                "approvedStart": roc_date(values[16]),
                "approvedEnd": roc_date(values[17]),
                "nextApplication": roc_date(values[18]),
            }
        )
        row += 1

    forms = {}
    form_positions = {}
    for sheet in workbook.worksheets[1:]:
        for rows in sheet.iter_rows(min_row=1, max_row=120, max_col=44):
            for cell in rows:
                water_right_no = clean_water_right(cell.value)
                if not water_right_no:
                    continue
                base_row = cell.row - 3
                base_col = cell.column - 2
                x = clean_number(sheet.cell(base_row + 11, base_col + 4).value)
                y = clean_number(sheet.cell(base_row + 11, base_col + 9).value)
                form = {
                    "completionDate": clean_text(sheet.cell(base_row + 8, base_col + 9).value),
                    "twd97X": compact_number(x),
                    "twd97Y": compact_number(y),
                }
                if is_valid_twd97(x, y):
                    form["latitude"], form["longitude"] = twd97_to_wgs84(x, y)
                forms[water_right_no] = form
                form_positions[(sheet.title, base_row, base_col)] = water_right_no

    photos = {record["waterRightNo"]: [] for record in registry}
    for sheet in workbook.worksheets[1:]:
        sheet_forms = [
            (base_row, base_col, water_right_no)
            for (sheet_name, base_row, base_col), water_right_no in form_positions.items()
            if sheet_name == sheet.title
        ]
        for image in sheet._images:
            anchor = getattr(image.anchor, "_from", None)
            if not anchor or image.width < 150 or image.height < 150:
                continue
            candidates = [
                item
                for item in sheet_forms
                if item[0] + 11 <= anchor.row <= item[0] + 15
                and item[1] <= anchor.col <= item[1] + 2
            ]
            if not candidates:
                continue
            _, _, water_right_no = min(
                candidates, key=lambda item: abs(anchor.row - (item[0] + 12)) + abs(anchor.col - item[1])
            )
            photos.setdefault(water_right_no, []).append(
                {"data": image._data(), "extension": image_extension(image)}
            )
    return registry, forms, photos


def file_bytes(path):
    try:
        return path.read_bytes()
    except OSError:
        return None


def unique_images(images):
    seen = set()
    result = []
    for image in images:
        digest = hashlib.sha256(image["data"]).hexdigest()
        if digest in seen:
            continue
        seen.add(digest)
        result.append(image)
    return result


def existing_photo_data(well):
    result = []
    for photo in well.get("photos") or []:
        stored_name = photo.get("storedName") or ""
        data = file_bytes(ATTACHMENTS_DIR / stored_name)
        if not data:
            continue
        extension = Path(stored_name).suffix.lower() or mimetypes.guess_extension(photo.get("mimeType") or "") or ".jpg"
        result.append({"data": data, "extension": extension})
    return result


def write_photos(well, new_images, timestamp, apply_changes):
    old_images = existing_photo_data(well)
    combined = unique_images(new_images)
    if len(combined) < len(old_images):
        combined = unique_images(combined + old_images)
    if not combined:
        return well.get("photos") or []
    if not apply_changes:
        return well.get("photos") or []

    station = well["station"]
    number = well["waterRightNo"]
    photos = []
    for index, image in enumerate(combined, 1):
        extension = image["extension"] if image["extension"].startswith(".") else f".{image['extension']}"
        stored_name = f"{station}-{number}-site-photo-{index}{extension.lower()}"
        target = ATTACHMENTS_DIR / stored_name
        target.write_bytes(image["data"])
        mime_type = mimetypes.guess_type(target.name)[0] or "image/jpeg"
        photos.append(
            {
                "id": f"photo-{station}-{number.lower()}-{index}",
                "name": f"{well['name']}現場照片{index}{target.suffix.lower()}",
                "mimeType": mime_type,
                "size": target.stat().st_size,
                "storedName": stored_name,
                "uploadedAt": timestamp,
            }
        )
    return photos


def find_water_right_pdf(root, station, name):
    if not root:
        return None
    station_dir = root / station
    if not station_dir.exists():
        return None
    exact = station_dir / f"{name}.pdf"
    if exact.exists():
        return exact
    matches = list(station_dir.glob(f"*{name}*.pdf"))
    return matches[0] if matches else None


def update_history(current_numbers):
    path = ROOT / "docs" / "data" / "pumping-history.json"
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    payload["records"] = [record for record in payload["records"] if record["waterRightNo"] in current_numbers]
    rights = sorted({record["waterRightNo"] for record in payload["records"]})
    payload["waterRightCount"] = len(rights)
    payload["wellCount"] = len(current_numbers)
    payload["recordCount"] = len(payload["records"])
    payload["monthlyRecordCount"] = len(payload["records"]) * 12
    payload["authorityCounts"] = {
        "臺中市政府": len({number for number in rights if number.startswith("B")}),
        "苗栗縣政府": len({number for number in rights if number.startswith("K")}),
    }
    payload["emptyWaterRightNos"] = sorted(current_numbers - set(rights))
    payload["anomalyRecordCount"] = sum(len(record.get("anomalies") or []) for record in payload["records"])
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    public_path = ROOT / "public" / "data" / "pumping-history.json"
    public_path.parent.mkdir(parents=True, exist_ok=True)
    public_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def update_current_pumping(wells, timestamp):
    path = ROOT / "docs" / "data" / "pumping-records" / "pumping-records-115.json"
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    existing = {record["waterRightNo"]: record for record in payload["records"]}
    month_keys = [
        "january", "february", "march", "april", "may", "june",
        "july", "august", "september", "october", "november", "december",
    ]
    records = []
    for index, well in enumerate(wells, 1):
        record = deepcopy(existing.get(well["waterRightNo"]))
        if not record:
            record = {
                "monthlyActualM3": {month: None for month in month_keys},
                "annualActualM3": None,
                "note": "新版井籍尚無抽水紀錄",
            }
        record.update(
            {
                "recordNo": index,
                "station": well["station"],
                "wellName": well["name"],
                "waterRightNo": well["waterRightNo"],
            }
        )
        records.append(record)
    payload["records"] = records
    payload["recordCount"] = len(records)
    payload["synchronizedAt"] = timestamp
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--water-right-root", type=Path)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()

    old_wells = json.loads(WELLS_PATH.read_text(encoding="utf-8-sig"))
    old_lookup = {well["waterRightNo"]: well for well in old_wells}
    registry, forms, workbook_photos = parse_workbook(args.workbook)
    timestamp = datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")
    source_name = args.workbook.name
    new_wells = []
    added = []
    changed = []

    for source in registry:
        number = source["waterRightNo"]
        old = deepcopy(old_lookup.get(number) or {})
        old_snapshot = deepcopy(old)
        is_new = not old
        if is_new:
            added.append(number)
        well = old or {
            "id": f"well-{source['station']}-{number.lower()}",
            "wellNumber": number,
            "district": "",
            "section": "",
            "startedAt": "",
            "status": "使用中",
            "publicNote": "",
            "internalNote": f"由 {source_name} 匯入。",
            "isPublic": True,
            "attachments": [],
            "photos": [],
            "auditTrail": [],
            "createdAt": timestamp,
            "createdBy": "system",
        }
        original = json.dumps(well, ensure_ascii=False, sort_keys=True)
        for field in (
            "station", "name", "electricityNo", "agriculturalPower", "address",
            "depthMeters", "diameterMm", "pumpHorsepower", "pumpOutletInch",
            "planFlowCms", "benefitedAreaHa", "registeredFlowCms", "irrigationSystem", "purpose",
        ):
            if source.get(field) is not None:
                well[field] = source[field]
        well["waterRightNo"] = number
        well["wellNumber"] = number
        well["managementUnit"] = f"{source['station']}工作站"
        if source.get("approvedStart") and source.get("approvedEnd"):
            well["waterRightPeriod"] = f"{source['approvedStart']} 至 {source['approvedEnd']}"
        if source.get("nextApplication"):
            well["nextApplicationPeriod"] = source["nextApplication"]
        form = forms.get(number) or {}
        if form.get("completionDate"):
            well["completionDate"] = form["completionDate"]
        if is_valid_twd97(form.get("twd97X"), form.get("twd97Y")):
            well["twd97X"] = str(form["twd97X"])
            well["twd97Y"] = str(form["twd97Y"])
            well["latitude"] = form["latitude"]
            well["longitude"] = form["longitude"]
        well["photos"] = write_photos(
            well, workbook_photos.get(number) or [], timestamp, args.apply
        )
        if is_new and not well.get("attachments"):
            pdf = find_water_right_pdf(args.water_right_root, source["station"], source["name"])
            if pdf:
                well["attachments"] = [
                    {
                        "id": f"source-pdf-{source['station']}-{number.lower()}",
                        "name": pdf.name,
                        "mimeType": "application/pdf",
                        "size": pdf.stat().st_size,
                        "storedName": str(pdf),
                        "uploadedAt": timestamp,
                    }
                ]
        comparable = deepcopy(well)
        comparable.pop("auditTrail", None)
        comparable.pop("updatedAt", None)
        comparable.pop("updatedBy", None)
        old_comparable = deepcopy(old_snapshot)
        old_comparable.pop("auditTrail", None)
        old_comparable.pop("updatedAt", None)
        old_comparable.pop("updatedBy", None)
        if is_new or comparable != old_comparable:
            changed.append(number)
            well.setdefault("auditTrail", []).append(
                {"action": "synchronized-workbook-1150810", "actor": "system", "at": timestamp}
            )
            well["updatedAt"] = timestamp
            well["updatedBy"] = "system"
        elif original:
            well = old_snapshot
        new_wells.append(well)

    removed = sorted(set(old_lookup) - {well["waterRightNo"] for well in new_wells})
    summary = {
        "source": source_name,
        "wells": len(new_wells),
        "forms": len(forms),
        "photosInWorkbook": sum(len(items) for items in workbook_photos.values()),
        "added": added,
        "removed": removed,
        "changed": changed,
    }
    if not args.apply:
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        return

    if len(new_wells) != 111 or len({well["waterRightNo"] for well in new_wells}) != 111:
        raise ValueError("新版井籍必須為111口且水權狀號不可重複")
    ATTACHMENTS_DIR.mkdir(parents=True, exist_ok=True)
    WELLS_PATH.write_text(json.dumps(new_wells, ensure_ascii=False, indent=2), encoding="utf-8")
    history = update_history({well["waterRightNo"] for well in new_wells})
    update_current_pumping(new_wells, timestamp)
    summary["historyWaterRights"] = history["waterRightCount"]
    summary["historyAnnualRecords"] = history["recordCount"]
    summary["historyEmpty"] = history["emptyWaterRightNos"]
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
